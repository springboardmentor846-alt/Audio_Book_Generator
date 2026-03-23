from pathlib import Path
from typing import IO


def _clean_text(raw: str) -> str:
    """Remove form-feeds, soft-hyphens, collapse consecutive blank lines."""
    lines = raw.replace("\f", "\n").replace("\u00ad", "").splitlines()
    out, prev_blank = [], False
    for line in lines:
        line = line.rstrip()
        blank = not line.strip()
        if blank and prev_blank:
            continue
        out.append(line)
        prev_blank = blank
    return "\n".join(out).strip()


def _extract_pdf(file_obj: IO[bytes]) -> str:
    """Try pdfplumber first, fall back to PyPDF2."""

    # --- pdfplumber ---
    try:
        import pdfplumber
        file_obj.seek(0)
        pages = []
        with pdfplumber.open(file_obj) as pdf:
            for page in pdf.pages:
                try:
                    pages.append(_clean_text(page.extract_text(x_tolerance=2, y_tolerance=3) or ""))
                except Exception:
                    pages.append("")
        text = "\n\n".join(p for p in pages if p)
        if text.strip():
            return text
    except Exception:
        pass

    # --- PyPDF2 fallback ---
    try:
        try:
            # Modern package name (recommended)
            from pypdf import PdfReader  # type: ignore
        except Exception:
            # Backwards compatibility
            from PyPDF2 import PdfReader  # type: ignore
        file_obj.seek(0)
        reader = PdfReader(file_obj)
        pages = []
        for page in reader.pages:
            try:
                pages.append(_clean_text(page.extract_text() or ""))
            except Exception:
                pages.append("")
        return "\n\n".join(p for p in pages if p)
    except Exception:
        return ""


def _extract_docx(file_obj: IO[bytes]) -> str:
    """Extract paragraphs and table cells from a Word document."""
    from docx import Document

    file_obj.seek(0)
    doc = Document(file_obj)
    parts = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    return "\n".join(parts)


_ENCODINGS = ("utf-8-sig", "utf-8", "latin-1", "cp1252")

def _extract_txt(file_obj: IO[bytes]) -> str:
    """Decode a plain-text file trying common encodings."""
    file_obj.seek(0)
    raw = file_obj.read()
    if isinstance(raw, str):
        return raw
    for enc in _ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def get_word_count(text: str) -> int:
    return len(text.split())


def extract_text(uploaded_file, **kwargs) -> str:
    """
    Extract text from a PDF, DOCX, or TXT uploaded file.
    Each format is handled independently — a missing library for one
    format will never affect extraction for another.
    """
    name: str = getattr(uploaded_file, "name", "") or ""
    suffix: str = Path(name).suffix.lower()

    if suffix == ".txt":
        return _extract_txt(uploaded_file)

    if suffix == ".docx":
        return _extract_docx(uploaded_file)

    if suffix == ".pdf":
        return _extract_pdf(uploaded_file)

    raise ValueError(
        f"Unsupported file type '{suffix}'. Please upload a PDF, DOCX, or TXT file."
    )